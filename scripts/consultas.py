"""Consultas estadísticas sobre el SAE con todos los criterios acordados aplicados.

Cada suceso queda con: DJP (Comuna, como en los informes), barrio oficial (nombre de la web), cuadrante,
cómo se asignó el barrio, si el barrio es de los incorporados el 26/07/2026, si estaba dentro de la zona
vigente en su fecha, y tiempos de despacho y arribo.

Uso (desde scripts/):
    python consultas.py --desde 2026-09-01 --hasta 2026-09-30 --por barrio
    python consultas.py --desde 2026-09 --por tipo --filtro "djp=15" --top 10
    python consultas.py --desde 2026-09 --por djp --comparar 2026-08 --tiempos
    python consultas.py --desde 2026-09 --por barrio --por categoria --excel ../informes/robos.xlsx --filtro "categoria=ROBO"
Dimensiones: djp, barrio, cuadrante, tipo, categoria, mes, semana, dia, dia_semana, hora, resultado, barrio_nuevo, en_zona.
Filtros (se pueden repetir): campo=valor (exacto, sin distinguir mayúsculas), campo~texto (contiene), positivo=1.
"""
import argparse
import calendar
import collections
import csv
import datetime as dt
import os
import statistics
import sys

import cobertura
import sae
from generar_informe import csv_mas_nuevo

BASE = os.path.dirname(os.path.abspath(__file__))
CAT = os.path.join(BASE, '..', 'catalogo')
DIAS = ['lunes', 'martes', 'miércoles', 'jueves', 'viernes', 'sábado', 'domingo']
csv.field_size_limit(10**9)


def _segundos(v):
    try:
        h, m, s = v.strip().split(':')
        return int(h) * 3600 + int(m) * 60 + int(float(s))
    except ValueError:
        return None


def _periodo(desde, hasta):
    if len(desde) == 7 and not hasta:
        y, m = map(int, desde.split('-'))
        return dt.date(y, m, 1), dt.date(y, m, calendar.monthrange(y, m)[1])
    return dt.date.fromisoformat(desde), dt.date.fromisoformat(hasta or desde)


def _catalogo():
    filas = list(csv.reader(open(os.path.join(CAT, 'catalogo_barrios.csv'), encoding='utf-8-sig'), delimiter=';'))[1:]
    excel_a_web = {f[3]: f[0] for f in filas}
    cuadrante = {f[0]: f[1] for f in filas}
    kml_a_web, por_coordenada = {}, set()
    for f in filas:
        for k in f[4].split(' / '):
            if k:
                kml_a_web[k] = f[0]
                if f[8].startswith('coordenada'):
                    por_coordenada.add(k)
    sis_a_web = {}
    for r in list(csv.reader(open(os.path.join(CAT, 'equivalencias_barrios.csv'), encoding='utf-8-sig'), delimiter=';'))[1:]:
        if r[1] in excel_a_web:
            sis_a_web[r[0]] = excel_a_web[r[1]]
    return sis_a_web, kml_a_web, por_coordenada, cuadrante


def cargar(path, desde, hasta):
    sis_a_web, kml_a_web, por_coord, cuadrante = _catalogo()
    cob = cobertura.Cobertura()
    out, vistos = [], set()
    with open(path, encoding='utf-8', errors='replace', newline='') as f:
        r = csv.reader(f)
        h = [sae._fix(x, True).strip() for x in next(r)]
        ix = {k: i for i, k in enumerate(h)}
        g = lambda x, c: x[ix[c]].strip() if c in ix else ''
        for x in r:
            if len(x) != len(h):
                continue
            try:
                fc = sae._fecha(g(x, 'Fecha Creacion'))
            except (ValueError, IndexError):
                continue
            if not (desde <= fc.date() <= hasta) or g(x, 'Suceso ID') in vistos:
                continue
            vistos.add(g(x, 'Suceso ID'))
            comuna = g(x, 'Comuna').replace('DISTRITO', '').strip()
            cuad = g(x, 'Cuadricula')
            djp = sae.FUERA if not cuad or not comuna else comuna
            if djp not in sae.DJP_INFORME and djp != sae.FUERA:
                continue                      # mismo criterio que los informes
            try:
                lat, lon = map(float, g(x, 'lat.long').split(','))
            except ValueError:
                lat = lon = None
            en_kml = cob.barrio_en(lat, lon) if lat is not None else None
            # Barrio: los 6 barrios sin cuadrícula propia van por coordenada; el resto por la cuadrícula del sistema
            if en_kml and en_kml[0] in por_coord:
                barrio, metodo = kml_a_web[en_kml[0]], 'coordenada'
            elif cuad and cuad in sis_a_web:
                barrio, metodo = sis_a_web[cuad], 'cuadrícula'
            elif en_kml:
                barrio, metodo = kml_a_web.get(en_kml[0], en_kml[0]), 'coordenada'
            else:
                barrio, metodo = '(fuera de los 234)' if lat is not None else '(sin dato)', 'sin dato' if lat is None else 'fuera'
            tipo = sae._fix(g(x, 'Tipo y Subtipo'), True)
            out.append({
                'fecha': fc, 'mes': fc.strftime('%Y-%m'), 'dia': fc.strftime('%Y-%m-%d'),
                'semana': (fc.date() - dt.timedelta(days=(fc.weekday() - 4) % 7)).strftime('vie %d/%m/%Y'),
                'dia_semana': DIAS[fc.weekday()], 'hora': fc.hour, 'djp': djp,
                'barrio': barrio, 'barrio_metodo': metodo, 'cuadrante': cuadrante.get(barrio, ''),
                'barrio_nuevo': 'sí' if en_kml and en_kml[1] else 'no',
                'en_zona': 'sí' if cob.cubierto(lat, lon, fc.date()) else 'no',
                'tipo': tipo, 'categoria': tipo.split(' (')[0], 'resultado': g(x, 'Resultado'),
                'positivo': g(x, 'Resultado') == 'Positivo', 'lat': lat, 'lon': lon,
                'aprehendidos': sae._num(g(x, 'aprehendidos totales')), 'mayores': sae._num(g(x, 'mayores aprehendidos')),
                'menores': sae._num(g(x, 'menores aprehendidos')),
                'secuestros': sum(sae._num(g(x, c)) for _, c in sae.SECUESTROS),
                't_despacho': _segundos(g(x, 'Tiempo Creación-Despacho')), 't_arribo': _segundos(g(x, 'Tiempo Despacho-Arribo')),
            })
    return out


def filtrar(suc, filtros):
    for fl in filtros:
        if '~' in fl:
            k, v = fl.split('~', 1)
            suc = [s for s in suc if v.lower() in str(s.get(k.strip(), '')).lower()]
        else:
            k, v = fl.split('=', 1)
            k, v = k.strip(), v.strip()
            if k == 'positivo':
                suc = [s for s in suc if s['positivo'] == (v in ('1', 'si', 'sí', 'true'))]
            else:
                suc = [s for s in suc if str(s.get(k, '')).lower() == v.lower()]
    return suc


def tabla(suc, dims, tiempos=False):
    grupos = collections.defaultdict(list)
    for s in suc:
        grupos[tuple(s[d] for d in dims)].append(s)
    filas = []
    for k, g in grupos.items():
        a, p = len(g), sum(s['positivo'] for s in g)
        f = dict(zip(dims, k), alertas=a, positivos=p, pct_positivos=round(p * 100 / a, 1) if a else 0,
                 aprehendidos=sum(s['aprehendidos'] for s in g), secuestros=sum(s['secuestros'] for s in g))
        if tiempos:
            td = [s['t_despacho'] for s in g if s['t_despacho'] is not None]
            ta = [s['t_arribo'] for s in g if s['t_arribo'] is not None]
            f['despacho_mediana_min'] = round(statistics.median(td) / 60, 1) if td else None
            f['arribo_mediana_min'] = round(statistics.median(ta) / 60, 1) if ta else None
            f['con_arribo_cargado'] = len(ta)
        filas.append(f)
    return filas


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--desde', required=True)
    ap.add_argument('--hasta')
    ap.add_argument('--por', action='append', default=[])
    ap.add_argument('--filtro', action='append', default=[])
    ap.add_argument('--comparar', nargs='+', metavar=('DESDE', 'HASTA'))
    ap.add_argument('--tiempos', action='store_true')
    ap.add_argument('--orden', default='alertas')
    ap.add_argument('--top', type=int)
    ap.add_argument('--excel')
    ap.add_argument('--csv')
    a = ap.parse_args()
    path = os.path.abspath(a.csv) if a.csv else csv_mas_nuevo()
    d1, d2 = _periodo(a.desde, a.hasta)
    dims = a.por or ['djp']
    suc = filtrar(cargar(path, d1, d2), a.filtro)
    filas = tabla(suc, dims, a.tiempos)
    if a.comparar:
        c1, c2 = _periodo(a.comparar[0], a.comparar[1] if len(a.comparar) > 1 else None)
        prev = {tuple(f[d] for d in dims): f for f in tabla(filtrar(cargar(path, c1, c2), a.filtro), dims)}
        for f in filas:
            q = prev.get(tuple(f[d] for d in dims), {})
            for m in ('alertas', 'positivos'):
                f[f'{m}_comparado'] = q.get(m, 0)
                f[f'{m}_variacion_%'] = round((f[m] - q.get(m, 0)) * 100 / q[m], 1) if q.get(m) else None
    num = lambda v: v if isinstance(v, (int, float)) and v is not None else -1
    filas.sort(key=lambda f: (-num(f.get(a.orden, 0)), str([f[d] for d in dims])))
    if a.top:
        filas = filas[:a.top]
    tot_a, tot_p = len(suc), sum(s['positivo'] for s in suc)
    print(f'Período {d1:%d/%m/%Y} al {d2:%d/%m/%Y} | filtros: {a.filtro or "ninguno"} | {tot_a} alertas, {tot_p} positivos | CSV: {os.path.basename(path)}')
    if filas:
        cols = list(filas[0].keys())
        anchos = {c: max(len(c), *(len(str(f[c])) for f in filas)) for c in cols}
        print('  '.join(c.ljust(anchos[c]) for c in cols))
        for f in filas:
            print('  '.join(str(f[c]).ljust(anchos[c]) for c in cols))
    if a.excel:
        exportar(a.excel, filas, suc, f'{d1:%d/%m/%Y} al {d2:%d/%m/%Y}', a.filtro)


def exportar(path, filas, suc, periodo, filtros):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = 'Resumen'
    ws.append([f'SAE · Cordobeses en Alerta — período {periodo}' + (f' — filtros: {", ".join(filtros)}' if filtros else '')])
    ws['A1'].font = Font(bold=True, size=12)
    if filas:
        ws.append(list(filas[0].keys()))
        for c in ws[2]:
            c.font = Font(bold=True, color='FFFFFF')
            c.fill = PatternFill('solid', fgColor='1F3A5F')
        for f in filas:
            ws.append(list(f.values()))
    ws2 = wb.create_sheet('Sucesos')
    campos = ['fecha', 'djp', 'barrio', 'cuadrante', 'barrio_metodo', 'barrio_nuevo', 'en_zona', 'tipo', 'resultado',
              'aprehendidos', 'mayores', 'menores', 'secuestros', 't_despacho', 't_arribo', 'lat', 'lon']
    ws2.append(campos)
    for c in ws2[1]:
        c.font = Font(bold=True)
    for s in suc:
        ws2.append([s[k] for k in campos])
    for w in (ws, ws2):
        for col in w.columns:
            w.column_dimensions[col[0].column_letter].width = min(45, max(10, *(len(str(c.value or '')) for c in col[:200])) + 2)
    wb.save(path)
    print(f'Excel: {os.path.abspath(path)}')


if __name__ == '__main__':
    main()
